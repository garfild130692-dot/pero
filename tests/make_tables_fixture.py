# Генерирует tests/fixtures/tables.docx средствами python-docx: ширины столбцов, границы, заливка,
# вертикальное выравнивание, высота строк, повтор заголовка, положение таблицы, границы из стиля.
# Запуск: python tests/make_tables_fixture.py tests/fixtures/tables.docx
import sys
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm


def el(tag, **attrs):
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(qn('w:' + k), str(v))
    return e


def set_widths(table, widths_cm):
    table.autofit = False
    for i, w in enumerate(widths_cm):
        table.columns[i].width = Cm(w)
        for cell in table.columns[i].cells:
            cell.width = Cm(w)


d = Document()

# 1) «Сетка таблицы»: ширины 3 / 5 / 8 см, заголовок повторяется, строка 2 — точно 1 см, по центру
d.add_paragraph('Таблица 1')
t = d.add_table(rows=3, cols=3, style='Table Grid')
set_widths(t, [3, 5, 8])
t.alignment = WD_TABLE_ALIGNMENT.CENTER
for r in range(3):
    for c in range(3):
        t.cell(r, c).text = f'A{r}{c}'
hdr = t.rows[0]._tr.get_or_add_trPr()
hdr.append(el('w:tblHeader'))
t.rows[1].height = Cm(1)
t.rows[1].height_rule = WD_ROW_HEIGHT_RULE.EXACTLY
tcPr = t.cell(0, 0)._tc.get_or_add_tcPr()
tcPr.append(el('w:shd', val='clear', color='auto', fill='DCEFE9'))
t.cell(1, 1).vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
t.cell(1, 2).vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.BOTTOM
b = el('w:tcBorders')
b.append(el('w:bottom', val='single', sz=24, space=0, color='C00000'))
t.cell(2, 2)._tc.get_or_add_tcPr().append(b)

# 2) Без стиля и без границ — в Word границ нет
d.add_paragraph('Таблица 2')
t2 = d.add_table(rows=2, cols=2)
for r in range(2):
    for c in range(2):
        t2.cell(r, c).text = f'B{r}{c}'

# 3) Прямые границы tblBorders: двойная рамка, внутри — штриховые
d.add_paragraph('Таблица 3')
t3 = d.add_table(rows=2, cols=2)
for r in range(2):
    for c in range(2):
        t3.cell(r, c).text = f'C{r}{c}'
tb = el('w:tblBorders')
for side in ('top', 'left', 'bottom', 'right'):
    tb.append(el('w:' + side, val='double', sz=4, space=0, color='1F6F5C'))
for side in ('insideH', 'insideV'):
    tb.append(el('w:' + side, val='dashed', sz=8, space=0, color='auto'))
t3._tbl.tblPr.append(tb)

# 4) Границы из стиля таблицы
d.add_paragraph('Таблица 4')
t4 = d.add_table(rows=2, cols=2, style='Light Grid Accent 1')
for r in range(2):
    for c in range(2):
        t4.cell(r, c).text = f'D{r}{c}'
d.save(sys.argv[1])
