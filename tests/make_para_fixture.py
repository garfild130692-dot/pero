# Генерирует tests/fixtures/para.docx средствами python-docx: интервалы, межстрочный интервал, отступы.
# Запуск: python tests/make_para_fixture.py tests/fixtures/para.docx
import sys
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_LINE_SPACING
from docx.shared import Cm, Pt

d = Document()
# «Обычный» как в документах по ГОСТ: первая строка 1,25 см, полуторный интервал, без интервала после
nf = d.styles['Normal'].paragraph_format
nf.first_line_indent = Cm(1.25)
nf.line_spacing = 1.5
nf.space_after = Pt(0)

d.add_heading('Абзацы', 1)
d.add_paragraph('Обычный абзац по умолчанию.')

p = d.add_paragraph('Перед 12, после 6.')
p.paragraph_format.space_before = Pt(12); p.paragraph_format.space_after = Pt(6)

p = d.add_paragraph('Одинарный.'); p.paragraph_format.line_spacing = 1.0
p = d.add_paragraph('Двойной.'); p.paragraph_format.line_spacing = 2.0
p = d.add_paragraph('Точно 18 pt.')
p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY; p.paragraph_format.line_spacing = Pt(18)
p = d.add_paragraph('Минимум 20 pt.')
p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.AT_LEAST; p.paragraph_format.line_spacing = Pt(20)

p = d.add_paragraph('Слева 2 см, справа 1 см.')
p.paragraph_format.left_indent = Cm(2); p.paragraph_format.right_indent = Cm(1)
p = d.add_paragraph('Выступ 1 см при отступе слева 1 см.')
p.paragraph_format.left_indent = Cm(1); p.paragraph_format.first_line_indent = Cm(-1)
p = d.add_paragraph('Без отступа первой строки.')
p.paragraph_format.first_line_indent = Cm(0)

quote = d.styles.add_style('Цитата Перо', WD_STYLE_TYPE.PARAGRAPH)
quote.base_style = d.styles['Normal']
quote.paragraph_format.left_indent = Cm(3)
quote.paragraph_format.space_before = Pt(6)
d.add_paragraph('Абзац пользовательского стиля.', style='Цитата Перо')

li = d.add_paragraph('Пункт списка с отступом', style='List Bullet')
li.paragraph_format.left_indent = Cm(4)

t = d.add_table(rows=1, cols=1)
t.cell(0, 0).text = 'Ячейка'
d.save(sys.argv[1])
