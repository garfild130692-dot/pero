# Генерирует tests/fixtures/fonts.docx средствами python-docx: шрифты, размеры, цвета, выделение, индексы.
# Запуск: python tests/make_fonts_fixture.py tests/fixtures/fonts.docx
import sys
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

d = Document()
# «Обычный» как в типовом русском документе — от него считаются умолчания документа
normal = d.styles['Normal']
normal.font.name = 'Times New Roman'
normal.font.size = Pt(14)

d.add_heading('Шрифты', 1)

p = d.add_paragraph('Обычный, ')
r = p.add_run('Arial 12'); r.font.name = 'Arial'; r.font.size = Pt(12)
p.add_run(', ')
r = p.add_run('красный'); r.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
p.add_run(', ')
r = p.add_run('жёлтый'); r.font.highlight_color = WD_COLOR_INDEX.YELLOW
p.add_run(', ')
r = p.add_run('заливка')
shd = OxmlElement('w:shd'); shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), 'FFC000')
r._r.get_or_add_rPr().append(shd)
p.add_run(', x')
r = p.add_run('2'); r.font.superscript = True
p.add_run(', H')
r = p.add_run('2'); r.font.subscript = True
p.add_run('O, ')
# явно указанные значения, совпадающие с умолчаниями, — не прямое форматирование
r = p.add_run('как обычный'); r.font.name = 'Times New Roman'; r.font.size = Pt(14)
p.add_run(', ')
r = p.add_run('10,5'); r.font.size = Pt(10.5)

code = d.styles.add_style('Код', WD_STYLE_TYPE.PARAGRAPH)
code.base_style = normal
code.font.name = 'Courier New'
code.font.size = Pt(10)
d.add_paragraph('print("Перо")', style='Код')

accent = d.styles.add_style('Акцент', WD_STYLE_TYPE.CHARACTER)
accent.font.bold = True
accent.font.color.rgb = RGBColor(0x1F, 0x6F, 0x5C)
p = d.add_paragraph('Символьный стиль: ')
p.add_run('акцент', style='Акцент')

# подчёркивание разных видов и цвета, прописные, малые прописные, межзнаковый интервал (этап 14)
from docx.enum.text import WD_UNDERLINE
p = d.add_paragraph('Вид: ')
r = p.add_run('двойное'); r.font.underline = WD_UNDERLINE.DOUBLE
r._r.get_or_add_rPr().find(qn('w:u')).set(qn('w:color'), 'C00000')
p.add_run(', ')
r = p.add_run('волна'); r.font.underline = WD_UNDERLINE.WAVY
p.add_run(', ')
r = p.add_run('прописные'); r.font.all_caps = True
p.add_run(', ')
r = p.add_run('малые'); r.font.small_caps = True
p.add_run(', ')
r = p.add_run('разреженный')
sp = OxmlElement('w:spacing'); sp.set(qn('w:val'), '40'); r._r.get_or_add_rPr().append(sp)
p.add_run(', ')
r = p.add_run('уплотнённый')
sp = OxmlElement('w:spacing'); sp.set(qn('w:val'), '-15'); r._r.get_or_add_rPr().append(sp)
d.save(sys.argv[1])
