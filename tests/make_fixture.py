# Генерирует тестовый docx средствами python-docx (эталонная «чужая» программа)
import sys
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
d = Document()
d.add_heading('Отчёт по продажам', 0)
d.add_heading('Раздел 1', 1)
p = d.add_paragraph('Обычный текст, ')
p.add_run('жирный').bold = True
p.add_run(', ')
p.add_run('курсив').italic = True
p.add_run(', ')
p.add_run('подчёркнутый').underline = True
p.add_run('\tпосле табуляции.')
c = d.add_paragraph('По центру'); c.alignment = WD_ALIGN_PARAGRAPH.CENTER
j = d.add_paragraph('По ширине'); j.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
d.add_paragraph('Пункт A', style='List Bullet')
d.add_paragraph('Пункт B', style='List Bullet')
d.add_paragraph('Вложенный B.1', style='List Bullet 2')
d.add_paragraph('Шаг 1', style='List Number')
d.add_paragraph('Шаг 2', style='List Number')
t = d.add_table(rows=3, cols=3)
for r in range(3):
    for k in range(3):
        t.cell(r, k).text = f'R{r}C{k}'
t.cell(0, 0).merge(t.cell(0, 1))
t.cell(1, 2).merge(t.cell(2, 2))
d.add_paragraph('Конец документа')
d.save(sys.argv[1])
