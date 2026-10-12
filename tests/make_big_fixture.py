# Генерирует tests/fixtures/big.docx — большой документ для замеров скорости (этап 17): ~300 страниц A4,
# заголовки трёх уровней, обычные абзацы с прямым форматированием, списки, таблицы (с шапкой, на несколько листов),
# рисунки. Содержимое детерминированное (без случайностей) — замеры сравнимы между запусками.
# Запуск: python tests/make_big_fixture.py
import io
import struct
import zlib

from docx import Document
from docx.shared import Cm, Pt

WORDS = ('договор поставка сторона оплата срок качество товар претензия уведомление гарантия реквизиты подпись '
         'исполнитель заказчик обязательство ответственность расчёт акт приёмка услуга работа период').split()


def png(w, h, seed):
    # простая картинка-градиент (без Pillow)
    rows = b''.join(bytes(1) + bytes(v for x in range(w) for v in ((x * 7 + seed * 31) % 256, (y * 5 + seed * 17) % 256,
                                                             (x + y + seed) % 256)) for y in range(h))
    chunk = lambda t, d: struct.pack('>I', len(d)) + t + d + struct.pack('>I', zlib.crc32(t + d) & 0xffffffff)
    return (b'\x89PNG\r\n\x1a\n' + chunk(b'IHDR', struct.pack('>IIBBBBB', w, h, 8, 2, 0, 0, 0))
            + chunk(b'IDAT', zlib.compress(rows)) + chunk(b'IEND', b''))


def text(n, k):
    return ' '.join(f'{WORDS[(k * 7 + i * 3) % len(WORDS)]}{"," if i % 11 == 10 else ""}' for i in range(n)).capitalize() + '.'


d = Document()
sec = d.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.left_margin, sec.right_margin, sec.top_margin, sec.bottom_margin = Cm(3), Cm(1.5), Cm(2), Cm(2)
d.styles['Normal'].font.name = 'Calibri'
d.styles['Normal'].font.size = Pt(11)

k = 0
for ch in range(1, 31):  # 30 глав по ~10 страниц
    d.add_heading(f'Глава {ch}. {text(4, ch)}', level=1)
    for sub in range(1, 5):
        d.add_heading(f'{ch}.{sub}. {text(5, ch * 10 + sub)}', level=2)
        for p in range(6):
            k += 1
            para = d.add_paragraph()
            para.add_run(text(40 + (k * 13) % 70, k))
            if k % 5 == 0:
                r = para.add_run(' Важно: ' + text(8, k + 1))
                r.bold = True
            if k % 7 == 0:
                r = para.add_run(' ' + text(6, k + 2))
                r.italic = True
        if sub == 2:
            for i in range(5):
                d.add_paragraph(text(10 + i * 3, k + i), style='List Bullet')
        if sub == 3:
            d.add_heading(f'{ch}.{sub}.1. {text(3, ch + sub)}', level=3)
            for i in range(4):
                d.add_paragraph(text(12 + i * 2, k + i * 5), style='List Number')
    # таблица на ~1,5 листа с шапкой
    rows = 25 + (ch % 3) * 10
    t = d.add_table(rows=rows, cols=5)
    t.style = 'Table Grid'
    for c, h in enumerate(('№', 'Наименование', 'Кол-во', 'Цена', 'Сумма')):
        t.cell(0, c).text = h
    for r in range(1, rows):
        vals = (str(r), text(3 + r % 5, r + ch), str(r * 3 % 17 + 1), f'{(r * 137 + ch) % 9000 + 100},00', f'{(r * 411 + ch) % 90000 + 1000},00')
        for c, v in enumerate(vals):
            t.cell(r, c).text = v
    t.rows[0].repeat = True
    tr = t.rows[0]._tr
    trPr = tr.get_or_add_trPr()
    from docx.oxml import OxmlElement
    hdr = OxmlElement('w:tblHeader')
    trPr.append(hdr)
    if ch % 3 == 0:
        d.add_picture(io.BytesIO(png(160, 100, ch)), width=Cm(8))
    d.add_paragraph(text(60, ch * 3))

d.save('tests/fixtures/big.docx')
print('ok')
