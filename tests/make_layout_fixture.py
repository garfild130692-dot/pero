# Генерирует tests/fixtures/layout.docx средствами python-docx: положение абзаца на странице, границы и заливка абзаца,
# позиции табуляции (выравнивание, заполнитель, стиль, удаление позиции стиля, шаг по умолчанию).
# Запуск: python tests/make_layout_fixture.py tests/fixtures/layout.docx
import sys
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.shared import Cm, Twips
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

d = Document()

# порядок элементов pPr по схеме — вставляем перед первым из «более поздних»
LATER_THAN_BDR = ['w:shd', 'w:tabs', 'w:spacing', 'w:ind', 'w:jc']
LATER_THAN_SHD = ['w:tabs', 'w:spacing', 'w:ind', 'w:jc']


def insert(pPr, el, later):
    for tag in later:
        a = pPr.find(qn(tag))
        if a is not None:
            a.addprevious(el)
            return
    pPr.append(el)


def border(p, sides, val='single', sz=4, space=1, color='auto'):
    bdr = OxmlElement('w:pBdr')
    for s in sides:
        e = OxmlElement(f'w:{s}')
        e.set(qn('w:val'), val); e.set(qn('w:sz'), str(sz)); e.set(qn('w:space'), str(space)); e.set(qn('w:color'), color)
        bdr.append(e)
    insert(p._p.get_or_add_pPr(), bdr, LATER_THAN_BDR)


def shade(p, fill):
    e = OxmlElement('w:shd'); e.set(qn('w:val'), 'clear'); e.set(qn('w:color'), 'auto'); e.set(qn('w:fill'), fill)
    insert(p._p.get_or_add_pPr(), e, LATER_THAN_SHD)


d.add_heading('Положение на странице', 1)
d.add_paragraph('Обычный абзац.')
p = d.add_paragraph('Не отрывать от следующего.'); p.paragraph_format.keep_with_next = True
p = d.add_paragraph('Не разрывать абзац.'); p.paragraph_format.keep_together = True
p = d.add_paragraph('С новой страницы.'); p.paragraph_format.page_break_before = True
p = d.add_paragraph('Без запрета висячих строк.'); p.paragraph_format.widow_control = False
h = d.add_heading('Заголовок, который можно отрывать', 2); h.paragraph_format.keep_with_next = False

# стиль со свойством «не отрывать» — значение переходит в абзац
st = d.styles.add_style('Подпись Перо', WD_STYLE_TYPE.PARAGRAPH)
st.base_style = d.styles['Normal']
st.paragraph_format.keep_with_next = True
d.add_paragraph('Абзац стиля с «не отрывать».', style='Подпись Перо')

p = d.add_paragraph('Рамка вокруг абзаца.'); border(p, ['top', 'left', 'bottom', 'right'], space=1)
p = d.add_paragraph('Нижняя граница двойная красная.'); border(p, ['bottom'], val='double', sz=6, color='C00000')
p = d.add_paragraph('Заливка жёлтым.'); shade(p, 'FFFF00')
p = d.add_paragraph('Рамка с заливкой, первый абзац группы.'); border(p, ['top', 'left', 'bottom', 'right', 'between']); shade(p, 'DEEBF7')
p = d.add_paragraph('Второй абзац группы.'); border(p, ['top', 'left', 'bottom', 'right', 'between']); shade(p, 'DEEBF7')

# ---------- Табуляция ----------
d.add_heading('Табуляция', 1)
p = d.add_paragraph('Москва\t10 октября 2026 г.')
p.paragraph_format.tab_stops.add_tab_stop(Twips(8630), WD_TAB_ALIGNMENT.RIGHT)
p = d.add_paragraph('Глава 1. Введение\t5')
p.paragraph_format.tab_stops.add_tab_stop(Twips(8630), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
p = d.add_paragraph('\tПо центру')
p.paragraph_format.tab_stops.add_tab_stop(Cm(8), WD_TAB_ALIGNMENT.CENTER)
p = d.add_paragraph('Итого:\t1234,56')
p.paragraph_format.tab_stops.add_tab_stop(Cm(10), WD_TAB_ALIGNMENT.DECIMAL)
d.add_paragraph('а\tб\tв\tг')  # позиции по умолчанию
p = d.add_paragraph('Слева 3 см:\tтекст\tдальше по умолчанию')
p.paragraph_format.tab_stops.add_tab_stop(Cm(3))
ts = d.styles.add_style('Реквизиты Перо', WD_STYLE_TYPE.PARAGRAPH)
ts.base_style = d.styles['Normal']
ts.paragraph_format.tab_stops.add_tab_stop(Cm(4))
ts.paragraph_format.tab_stops.add_tab_stop(Cm(9), WD_TAB_ALIGNMENT.RIGHT)
d.add_paragraph('Стиль:\tчетыре\tдевять', style='Реквизиты Перо')
p = d.add_paragraph('Стиль без 4 см:\tдевять', style='Реквизиты Перо')
p.paragraph_format.tab_stops.add_tab_stop(Cm(4), WD_TAB_ALIGNMENT.CLEAR)

d.save(sys.argv[1])
