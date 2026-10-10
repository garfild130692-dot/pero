# Генерирует tests/fixtures/pages.docx средствами python-docx — эталон разбиения на страницы (сверка с Word через COM):
# A4, поля как в русском Word, Calibri 11; абзацы разной длины, заголовки у конца страниц, «не разрывать»,
# абзац без запрета висячих строк, длинные таблицы с повтором шапки, разрыв страницы, «с новой страницы», список.
# Запуск: python tests/make_pages_fixture.py tests/fixtures/pages.docx
import random
import sys
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

random.seed(7)
WORDS = ('договор сторона обязательство исполнение порядок оплата срок поставка товар качество приемка документ '
         'уведомление ответственность штраф неустойка претензия спор арбитраж изменение дополнение соглашение '
         'реквизиты подпись печать организация директор представитель основание устав доверенность период '
         'стоимость налог счет акт сверка расчет платеж банк обеспечение гарантия риск убыток').split()


def text(n):
    w = [random.choice(WORDS) for _ in range(n)]
    w[0] = w[0].capitalize()
    out, i = [], 0
    while i < len(w):
        k = random.randint(6, 14)
        out.append(' '.join(w[i:i + k]) + ('.' if random.random() < .7 else ','))
        i += k
    s = ' '.join(out)
    return s[:-1] + '.'


d = Document()
sec = d.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.orientation = WD_ORIENT.PORTRAIT
sec.top_margin, sec.bottom_margin, sec.left_margin, sec.right_margin = Cm(2), Cm(2), Cm(3), Cm(1.5)
st = d.styles['Normal']
st.font.name = 'Calibri'
st.font.size = Pt(11)
rpr = st.element.get_or_add_rPr()
fonts = rpr.find(qn('w:rFonts'))
for a in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
    fonts.set(qn(a), 'Calibri')
pf = st.paragraph_format
pf.space_after = Pt(8)
pf.line_spacing = 1.08
for h in ('Heading 1', 'Heading 2'):
    hs = d.styles[h]
    hs.font.name = 'Calibri Light'
    hr = hs.element.get_or_add_rPr().find(qn('w:rFonts'))
    for a in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
        hr.set(qn(a), 'Calibri Light')
    for a in ('w:asciiTheme', 'w:hAnsiTheme', 'w:cstheme', 'w:eastAsiaTheme'):
        if hr.get(qn(a)) is not None:
            del hr.attrib[qn(a)]


def para(n, **kw):
    p = d.add_paragraph(text(n))
    p.alignment = kw.get('align', WD_ALIGN_PARAGRAPH.JUSTIFY)
    f = p.paragraph_format
    if kw.get('first'):
        f.first_line_indent = Cm(1.25)
    for k in ('keep_with_next', 'keep_together', 'page_break_before', 'widow_control'):
        if k in kw:
            setattr(f, k, kw[k])
    return p


def repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    trPr.append(OxmlElement('w:tblHeader'))


d.add_heading('Договор поставки № 1', 1)
for i in range(9):
    para(random.randint(40, 160), first=True)
d.add_heading('1. Предмет договора', 2)
for i in range(7):
    para(random.randint(30, 120), first=i % 2 == 0)
para(140, keep_together=True)                       # не разрывать
for i in range(5):
    para(random.randint(60, 140))
para(150, widow_control=False)                      # без запрета висячих строк
d.add_heading('2. Спецификация', 2)
t = d.add_table(rows=1, cols=4)
t.style = 'Table Grid'
for c, h in zip(t.rows[0].cells, ('№', 'Наименование', 'Кол-во', 'Цена')):
    c.text = h
repeat_header(t.rows[0])
for i in range(55):
    r = t.add_row().cells
    r[0].text = str(i + 1)
    r[1].text = text(random.randint(2, 14))
    r[2].text = str(random.randint(1, 500))
    r[3].text = f'{random.randint(100, 99999)},00'
for i in range(4):
    para(random.randint(40, 120))
p = d.add_paragraph('Перед разрывом страницы.')
p.add_run().add_break(WD_BREAK.PAGE)
d.add_heading('3. Порядок расчетов', 2)
for i in range(6):
    para(random.randint(50, 150), first=True)
for i in range(12):
    d.add_paragraph(text(random.randint(8, 30)), style='List Number')
para(60, page_break_before=True)                    # с новой страницы
for i in range(10):
    para(random.randint(30, 170))
d.add_heading('4. Ответственность сторон', 2)
for i in range(14):
    para(random.randint(20, 110))
t = d.add_table(rows=1, cols=3)
t.style = 'Table Grid'
for c, h in zip(t.rows[0].cells, ('Показатель', 'Значение', 'Примечание')):
    c.text = h
repeat_header(t.rows[0])
for i in range(40):
    r = t.add_row().cells
    r[0].text = text(random.randint(1, 4))
    r[1].text = str(random.randint(1, 1000))
    r[2].text = text(random.randint(1, 20)) if i % 3 else ''
for i in range(8):
    para(random.randint(40, 140))
d.save(sys.argv[1])
