# Генерирует tests/out/bigbook.xlsx — большая книга для замеров (этап 17): 100 000 строк × 10 столбцов: номер, дата,
# текст (общие строки), числа, формулы (произведение, ЕСЛИ, ВПР по справочнику, нарастающий итог по скользящему окну),
# итоги СУММ/СРЗНАЧ/СЧЁТЕСЛИ по целым столбцам. Файл большой — в репозиторий не кладётся.
# Запуск: python tests/make_bigbook_fixture.py [строк]
import datetime
import sys

from openpyxl import Workbook

N = int(sys.argv[1]) if len(sys.argv) > 1 else 100000
wb = Workbook(write_only=True)
ws = wb.create_sheet('Данные')
ws.append(['№', 'Дата', 'Товар', 'Кол-во', 'Цена', 'Сумма', 'Скидка', 'Склад', 'Группа', 'Итог 10'])
goods = [f'Товар {i}' for i in range(500)]
d0 = datetime.date(2024, 1, 1)
for r in range(2, N + 2):
    i = r - 2
    ws.append([i + 1, d0 + datetime.timedelta(days=i % 700), goods[(i * 7) % 500], (i * 13) % 50 + 1, round(((i * 37) % 9000) / 7 + 10, 2),
               f'=D{r}*E{r}', f'=IF(F{r}>5000,F{r}*0.05,0)', f'=VLOOKUP(C{r},Справочник!$A$1:$B$500,2,FALSE)',
               (i * 3) % 12, f'=SUM(F{max(2, r - 9)}:F{r})'])
sp = wb.create_sheet('Справочник')
for i, g in enumerate(goods):
    sp.append([g, f'Склад {i % 9 + 1}'])
it = wb.create_sheet('Итоги')
it.append(['Сумма', '=SUM(Данные!F:F)'])
it.append(['Среднее', '=AVERAGE(Данные!F:F)'])
it.append(['Крупных', '=COUNTIF(Данные!F:F,">5000")'])
wb.save('tests/out/bigbook.xlsx')
print('ok', N)
